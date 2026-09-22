from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(r"C:\Users\harsh\OneDrive\Documents\ChatGPT\SIH\output\documents\BhuNirnay_SIH_26019_Project_Blueprint.docx")

NAVY = "0F2D5C"
GREEN = "1A6B3C"
LIGHT_BLUE = "EAF1FA"
PALE_GREEN = "EAF5EE"
GRAY = "D9D9D9"
TEXT = "202124"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRAY):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_col_width(cell, width_cm):
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def add_page_field(paragraph):
    run = paragraph.add_run("Page ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def font(run, size=10.5, bold=False, color=TEXT, italic=False):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        font(r)
    else:
        r = p.add_run(text)
        font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Cm(0.6)
        r = p.add_run(item)
        font(r)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    font(r, size=15 if level == 1 else 12, bold=True, color=NAVY)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    for idx, label in enumerate(headers):
        cell = header_cells[idx]
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        font(r, size=8.6, bold=True, color="FFFFFF")
        if widths:
            set_col_width(cell, widths[idx])
    for row_num, row_values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cell = cells[idx]
            set_cell_border(cell)
            set_cell_margins(cell)
            set_cell_shading(cell, "FFFFFF" if row_num % 2 == 0 else "F6F9FC")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            font(r, size=8.5)
            if widths:
                set_col_width(cell, widths[idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def checkbox_table(doc, title, items):
    add_heading(doc, title, 2)
    rows = [["Done", item] for item in items]
    table = add_table(doc, ["", "Verification item"], rows, widths=[1.25, 15.2])
    for row in table.rows[1:]:
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(10.5)

    # Header and footer
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("BhuNirnay | SIH 26019 Project Blueprint")
    font(r, size=8.5, bold=True, color=GREEN)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Internal working document | Update after every approved milestone | ")
    font(r, size=8, color="666666")
    add_page_field(footer)

    # Cover page
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(52)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("BhuNirnay")
    font(r, size=28, bold=True, color=NAVY)
    p = doc.add_paragraph()
    r = p.add_run("SIH 26019 Project Requirements and Delivery Blueprint")
    font(r, size=18, bold=True, color=TEXT)
    p.paragraph_format.space_after = Pt(10)
    p = doc.add_paragraph()
    r = p.add_run("National Digital Platform for Research Policy Innovation and Evidence Based Land Governance")
    font(r, size=12, color=GREEN, italic=True)
    p.paragraph_format.space_after = Pt(22)
    add_body(doc, "Purpose: This is the team’s single working plan for the full project journey. It translates the SIH problem statement into buildable modules, defines what counts as complete, and records the order in which we will deliver and test every capability.")
    add_table(doc, ["Document control", "Value"], [
        ["Problem statement", "SIH 26019 | Ministry of Rural Development | Department of Land Resources PME Division"],
        ["Version", "1.0"],
        ["Prepared on", "20 September 2026"],
        ["Primary outcome", "A secure national knowledge and decision-support platform for evidence-based land governance"],
        ["Working rule", "No feature is called complete until it is implemented, tested, documented, and demonstrable."],
    ], widths=[4.2, 12.2])
    doc.add_page_break()

    add_heading(doc, "How to Use This Blueprint")
    add_body(doc, "This document is the project control sheet. At the start of each work session, select the next unchecked item from the current phase. At the end of the session, update its status, add evidence of testing, and record any decision or risk. Do not build features outside the agreed phase unless the team explicitly changes the plan.")
    add_heading(doc, "Project Vision")
    add_body(doc, "BhuNirnay will be a centralized and collaborative digital ecosystem where researchers, policymakers, government agencies, institutions and public users can discover trusted land-governance knowledge, analyse evidence, compare policy choices and access decision-support tools with appropriate permissions.")
    add_heading(doc, "Success Criteria")
    add_bullets(doc, [
        "Every AI answer is grounded in approved evidence, shows source title and page range, and refuses to invent information when evidence is missing.",
        "Users can discover policies, research papers, datasets, legal documents and case studies through filters, keyword search and semantic search.",
        "Recommendations explain why a resource, scheme, policy or dataset is relevant to the question and user context.",
        "GIS, analytics and policy-simulation results show their source, assumptions and limits.",
        "Role-based permissions protect administrative workflows and private documents.",
        "The final demo can be run reliably from the hosted website and has a repeatable local setup." 
    ])

    add_heading(doc, "Official Requirement Traceability", 1)
    add_body(doc, "The table below is the non-negotiable mapping from the SIH problem statement to BhuNirnay features. Each row must have a visible demo path before final submission.")
    requirements = [
        ["Central repository", "Research, policies, datasets, legal documents and case studies", "Supabase document catalog, private storage, metadata, review workflow", "Partial"],
        ["AI search and recommendation", "Discover relevant resources", "Hybrid search, RAG answers, related-resource recommendations", "Search partial; recommendations pending"],
        ["Collaborative workspaces", "Researchers, policymakers, institutions and agencies collaborate", "Projects, notes, comments, assignments, approval workflow", "UI demo only"],
        ["GIS visualization", "Land use, climate, infrastructure and policy impacts", "GIS layers, district/parcel views, sourced indicators", "UI prototype only"],
        ["Advanced analytics", "Evaluate policy effectiveness and emerging trends", "Indicator datasets, charts, trend calculations and source notes", "Prototype calculations only"],
        ["Policy simulation", "Assess reforms before implementation", "Transparent assumptions, scenario inputs, saved results", "Prototype formula only"],
        ["Data integration", "Satellite, remote sensing, land records, socio-economic and GIS data", "Source connectors and dataset ingestion pipeline", "Pending"],
        ["AI research tools", "Trend analysis, literature synthesis, predictive and scenario analysis", "RAG synthesis, evidence summaries and later analytical models", "RAG partial"],
        ["Innovation portal", "Hackathons, grants, pilot projects and competitions", "Challenge/project registry and application workflow", "Pending"],
        ["Dashboards", "Research, policy, land use, climate, disputes, projects and geospatial insights", "Role-specific dashboards backed by actual indicators", "UI prototype only"],
        ["Role based access", "Researchers, officials, institutions and public users", "Supabase auth, backend JWT validation, permissions and audit logs", "Schema partial"],
        ["Interoperability APIs", "Government platforms, research databases, GIS and governance systems", "Documented REST API, import/export contracts and adapters", "Backend base only"],
    ]
    add_table(doc, ["Requirement", "Official intent", "BhuNirnay implementation", "Status"], requirements, widths=[3.0, 4.1, 6.5, 2.8])

    add_heading(doc, "Users and Permissions")
    add_table(doc, ["Role", "Primary needs", "Allowed actions in final system"], [
        ["Public user", "Discover public resources and basic insights", "Search approved public metadata, view public dashboards and use safe AI queries"],
        ["Researcher", "Conduct studies and share evidence", "Create projects, upload sources for review, save research notes and access research workspace"],
        ["Government officer", "Evaluate policies and monitor indicators", "Use policy simulator, restricted datasets, officer dashboards and approved reports"],
        ["Institution manager", "Coordinate institutional contributions", "Manage affiliated projects and review institutional submissions"],
        ["Administrator", "Protect trust and operations", "Approve sources, manage roles, monitor audit logs and manage data lifecycle"],
    ], widths=[3.0, 5.3, 8.1])

    add_heading(doc, "Solution Architecture")
    add_body(doc, "The architecture separates public presentation from confidential processing. Secrets and privileged database actions remain on the backend; no Gemini or Supabase service key is ever placed in browser code or a public Vercel variable.")
    add_table(doc, ["Layer", "Technology", "Responsibility"], [
        ["Web application", "React and Vite on Vercel", "User interface, maps, dashboards, search forms and result display"],
        ["API service", "FastAPI on Render, Railway or government-compatible host", "Authentication checks, RAG orchestration, analytics APIs and audit events"],
        ["Knowledge database", "Supabase PostgreSQL and pgvector", "Document metadata, chunks, embeddings, roles, projects and indicators"],
        ["AI services", "Gemini embedding and generation APIs", "Embeddings, grounded answer synthesis and controlled research assistance"],
        ["Source storage", "Private Supabase Storage", "Original PDFs, datasets and approved source files"],
        ["Future search layer", "PostgreSQL full text plus vector retrieval", "Hybrid retrieval, filters, relevance ranking and recommendations"],
    ], widths=[3.0, 5.5, 7.9])

    add_heading(doc, "Repository and Data Governance")
    add_body(doc, "Every stored resource must be traceable. The platform must distinguish verified government evidence from unverified research notes and clearly disclose whether information is public, restricted or pending review.")
    add_bullets(doc, [
        "Accept only public, permitted or institution-authorized sources. Record source URL, organization, document type, state, district, publication date, tags, checksum and licence/reuse note.",
        "Upload original source files to private storage. The browser must never receive privileged storage access.",
        "Set every new document to pending review. An administrator verifies provenance and extracted text before approval.",
        "Store a checksum and version/review date so updated policies can be re-ingested and superseded safely.",
        "For scanned PDFs, run OCR and a human readability check before creating embeddings.",
        "Do not present synthetic or placeholder papers, schemes, metrics or citations as government evidence." 
    ])

    add_heading(doc, "AI Search and Recommendation Engine")
    add_body(doc, "This is the core intelligence feature. It has two distinct outputs: evidence-grounded answers and explainable recommendations. We will not describe a generic chatbot as the full engine.")
    add_heading(doc, "Search Process", 2)
    add_table(doc, ["Step", "What the system does", "Completion evidence"], [
        ["1 Query intake", "Validate question length, identity, role and optional filters", "Invalid request is rejected; user role comes from verified token, not browser text"],
        ["2 Query understanding", "Detect subject, location, time period, document type and user intent", "Search applies Maharashtra, district, year or topic filters when supplied"],
        ["3 Hybrid retrieval", "Run keyword/full-text and vector search over approved resources", "Exact policy codes and semantic paraphrases both retrieve relevant results"],
        ["4 Reranking", "Prioritize authority, topical relevance, state/district match, recency and source quality", "Returned result order is explainable"],
        ["5 Grounded generation", "Send only selected evidence to the model with strict citation instructions", "Every factual statement links to a visible source/page range"],
        ["6 Confidence handling", "Refuse or qualify answers with weak evidence", "No-evidence query returns a transparent insufficiency message"],
        ["7 Recommendations", "Suggest related policies, documents, datasets, schemes and case studies", "Each recommendation carries a plain-language reason and source"],
        ["8 Audit and feedback", "Store non-sensitive query metadata, selected sources and user rating", "Admin can inspect quality trends without exposing private content"],
    ], widths=[2.2, 8.2, 6.0])
    add_heading(doc, "Recommendation Rules", 2)
    add_body(doc, "The first recommendation engine will be rule-based and explainable. It will combine semantic relevance, topic match, state/district match, document authority and recency. Personalization will be added only after consent, role verification and privacy review.")
    add_table(doc, ["Signal", "Initial weight", "Explanation shown to user"], [
        ["Semantic relevance", "40%", "Matches the meaning of your question"],
        ["Location match", "25%", "Applies to your selected state or district"],
        ["Authority", "15%", "Published by an approved government or verified institution"],
        ["Recency", "10%", "More recent verified guidance is preferred"],
        ["Task and role relevance", "10%", "Useful for your selected research or policy task"],
    ], widths=[4.2, 3.2, 9.0])

    add_heading(doc, "RAG Source Ingestion Standard")
    add_table(doc, ["Stage", "Required action", "Gate"], [
        ["Discover", "Find official source and verify organization, URL and reuse permission", "Source log is completed"],
        ["Register", "Add title, type, tags, geography, date and provenance metadata", "Required metadata is present"],
        ["Extract", "Read PDF page by page; OCR scans where required", "Text is readable and page numbers are preserved"],
        ["Chunk", "Create coherent overlapping chunks without mixing unrelated pages", "Chunks are inspectable and citation-friendly"],
        ["Embed", "Generate vectors with rate-limit handling and checkpoint/resume", "All chunks have the expected vector dimension"],
        ["Review", "Administrator checks source, extraction and classification", "Document is explicitly approved"],
        ["Evaluate", "Ask at least three representative questions and verify sources", "Expected evidence is retrieved"],
        ["Maintain", "Record review date and replace superseded documents", "Outdated sources are archived, not silently overwritten"],
    ], widths=[3.0, 9.0, 4.4])

    add_heading(doc, "GIS Analytics and Policy Decision Support")
    add_bullets(doc, [
        "Phase 1: use clearly labelled sample district/parcel data only for interface validation.",
        "Phase 2: ingest official, reusable district-level indicators with source URL, date, geography and methodology.",
        "Phase 3: connect GIS layers for land use, climate vulnerability, infrastructure and policy context. Every map layer displays its source and update date.",
        "Phase 4: make policy simulation transparent. Inputs, formula assumptions, uncertainty and source data must be visible. It is decision support, never an official legal determination.",
        "Phase 5: add scenario comparison and exportable policy briefs with evidence citations." 
    ])

    add_heading(doc, "Collaborative Research and Innovation Features")
    add_table(doc, ["Module", "Minimum final capability", "Later enhancement"], [
        ["Research workspace", "Create a research project, invite permitted members, attach approved sources and save notes", "Task board, workflows and peer review"],
        ["Collaborative notes", "Author, timestamp, tags, source/parcel reference and moderation status", "Version history and threaded comments"],
        ["Innovation portal", "List challenges, pilots and research calls with criteria and deadlines", "Applications, reviewer workflow and project outcomes"],
        ["Reporting", "Export a research brief with selected evidence and assumptions", "Scheduled reports and institutional dashboards"],
    ], widths=[3.7, 8.0, 4.7])

    add_heading(doc, "Delivery Phases and Exit Gates")
    phases = [
        ["0 Foundation audit", "Verify repository, environment files, Supabase schema, source approval and deployment configuration", "One team setup guide and a known-good local run"],
        ["1 Live RAG", "Deploy backend, connect Vercel, verify one real source end-to-end", "Hosted question returns cited approved evidence"],
        ["2 Knowledge catalog", "Add source metadata screens and ingest first verified source batch", "At least 8 to 12 approved high-value sources"],
        ["3 Hybrid search", "Add keyword, vector, filter and reranking retrieval", "Search works for exact names and natural-language questions"],
        ["4 Recommendations", "Return related documents, policies, datasets and schemes with reasons", "Recommendations are sourced and explainable"],
        ["5 Real decision modules", "Replace priority mock data in analytics, GIS and simulations", "Every shown metric has a source and methodology"],
        ["6 Collaboration and roles", "JWT-backed access, project workflow, audit events and moderation", "Role abuse tests pass"],
        ["7 Evaluation and polish", "Benchmark AI, test UI, fix accessibility, rehearse demo and prepare PPT evidence", "Demo script passes twice without manual workaround"],
    ]
    add_table(doc, ["Phase", "Scope", "Exit gate"], phases, widths=[3.1, 8.2, 5.1])

    add_heading(doc, "Responsibilities")
    add_table(doc, ["Workstream", "Primary responsibility", "Acceptance output"], [
        ["Presentation and narrative", "Project lead / PPT owner", "Problem, architecture, impact, demo flow, metrics and limitations explained clearly"],
        ["Frontend experience", "Website teammate", "Responsive UI, forms, state handling, loading/errors and deployment-ready build"],
        ["Backend and AI", "Technical workstream", "Secure API, RAG, search, recommendation logic, ingestion and evaluation"],
        ["Data and verification", "Assigned data owner plus administrator", "Source log, licences, metadata, approval and quality checks"],
        ["GIS and analytics", "Assigned data/analytics owner", "Verified layers, indicators, methodology and limitations"],
        ["Quality and demo", "Whole team", "Test record, evidence screenshots and stable end-to-end rehearsal"],
    ], widths=[3.3, 6.2, 6.9])

    add_heading(doc, "Security and Deployment Rules")
    add_bullets(doc, [
        "Vercel hosts only the React frontend. The secure FastAPI backend is deployed separately.",
        "Never put GEMINI_API_KEY, SUPABASE_SECRET_KEY, passwords or service tokens in GitHub, frontend JavaScript or Vercel public variables.",
        "VITE_API_BASE_URL may be public because it is only the API address.",
        "Backend validates Supabase JWT and fetches user role from the database. Never trust a role supplied by the browser.",
        "Restrict CORS to the exact production domain and known preview domains, rather than all Vercel websites.",
        "Use rate limiting, request validation, structured logs and error messages that do not reveal secrets.",
        "Use private storage for original source files and issue time-limited access only after authorization." 
    ])

    add_heading(doc, "Testing and Evidence Checklist")
    checkbox_table(doc, "Before Calling Any Feature Complete", [
        "Feature has a clear user story and accepted scope.",
        "Happy path is manually tested on local environment.",
        "Invalid input, empty data, denied permission and API failure have understandable behaviour.",
        "No secrets appear in Git status, browser console, deployment logs or screenshots.",
        "Data shown on screen carries source, date and limitation where relevant.",
        "RAG answer test confirms the cited page actually supports the claim.",
        "Feature works on the hosted deployment, not only localhost.",
        "A team member other than the builder can follow the test steps." 
    ])
    checkbox_table(doc, "RAG Evaluation Set", [
        "Prepare 30 to 50 questions across national policy, Maharashtra land records, GIS and scheme discovery.",
        "For each question record expected source, expected page/range, core fact and whether refusal is expected.",
        "Measure retrieval relevance, citation correctness, answer groundedness, refusal quality and response time.",
        "Fix retrieval before changing prompts when the correct source is not retrieved.",
        "Keep test results as proof for the SIH jury and final PPT." 
    ])

    add_heading(doc, "Known Risks and Controls")
    add_table(doc, ["Risk", "Impact", "Control"], [
        ["Too little verified data", "Weak or incomplete answers", "Ingest sources by curated priority list; refuse when evidence is insufficient"],
        ["Rate limits on free AI tier", "Slow ingestion or temporary failure", "Single/batched requests, delay, exponential backoff, checkpoint/resume and optional offline fallback"],
        ["Mock data presented as real", "Loss of credibility", "Clearly label prototypes and replace them phase by phase with sourced datasets"],
        ["Secret exposed in frontend", "Database/AI compromise", "Backend-only secret storage, Git ignore, secret scan and key rotation procedure"],
        ["Weak citations or hallucinations", "Unsafe policy guidance", "Grounded prompt, evidence threshold, citation verification and benchmark tests"],
        ["Unstable public deployment", "Demo failure", "Health endpoint, monitoring, deployment checklist and local backup demo route"],
        ["Scope overload", "Incomplete final product", "Phase gates: complete core search/recommendation and evidence trust before optional features"],
    ], widths=[4.0, 4.5, 7.9])

    add_heading(doc, "Current Baseline and Immediate Next Actions")
    add_body(doc, "Current baseline as of 20 September 2026: the website interface exists; Supabase schema and basic role model exist; a FastAPI RAG backend exists; one official DoLR DILRMP 3.0 document has been ingested as 84 chunks and tested locally. The hosted frontend has deployment-compatible code, but a secure public backend deployment and end-to-end hosted test still need confirmation.")
    add_table(doc, ["Priority", "Next action", "Owner", "Proof of completion"], [
        ["P0", "Deploy FastAPI backend and set private backend environment variables", "Technical workstream", "Public /health endpoint returns healthy"],
        ["P0", "Set VITE_API_BASE_URL in Vercel and redeploy frontend", "Deployment owner", "Hosted website reaches API without fallback"],
        ["P0", "Run five live RAG questions and verify citations", "Whole team", "Test sheet with screenshot and sources"],
        ["P1", "Create verified source catalog and ingest first 8 to 12 sources", "Data owner plus administrator", "Approved documents and retrieval tests"],
        ["P1", "Implement hybrid search and filters", "Backend workstream", "Exact and semantic search tests pass"],
        ["P1", "Build related-resource recommendation cards", "Frontend and backend", "Recommendations show explanation and source"],
    ], widths=[1.4, 7.5, 3.0, 4.5])

    add_heading(doc, "Decision Log Template")
    add_body(doc, "Add a new row whenever the team makes an important technical, data or scope decision. This prevents repeated discussion and protects the final presentation from contradictions.")
    add_table(doc, ["Date", "Decision", "Reason", "Owner", "Effect on plan"], [
        ["20 Sep 2026", "Use RAG rather than train a model from scratch", "Lower cost, source citations and faster policy updates", "Team", "RAG becomes the primary AI-search approach"],
        ["20 Sep 2026", "Keep AI and Supabase service secrets on backend only", "Protect credentials and administrative database access", "Team", "Deploy FastAPI separately from Vercel frontend"],
    ], widths=[2.3, 4.3, 5.0, 2.0, 2.8])

    add_heading(doc, "Final Demonstration Flow")
    add_bullets(doc, [
        "Open the hosted BhuNirnay platform and identify the user role.",
        "Search a policy question and show cited answer from an approved government document.",
        "Apply a state, document type or topic filter and show retrieval changes.",
        "Open an explainable recommendation for a related policy, dataset or scheme.",
        "Show a GIS/analytics view with source/date disclosure.",
        "Run a policy scenario, show assumptions and compare outcomes.",
        "Show administrator source approval and audit trail in a prepared safe demo account.",
        "Close with metrics: sources indexed, questions tested, citation accuracy and expected governance impact." 
    ])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
